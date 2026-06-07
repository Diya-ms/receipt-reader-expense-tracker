"""
File: backend/routes/export.py
Export expenses to CSV, Excel, and PDF.
"""
import io
import csv
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file, make_response
from flask_jwt_extended import jwt_required, get_jwt_identity

from models.expense import Expense

export_bp = Blueprint('export', __name__)


def _get_user_id():
    return int(get_jwt_identity())


def _get_filtered_expenses(uid):
    q = Expense.query.filter_by(user_id=uid)
    if request.args.get('date_from'):
        q = q.filter(Expense.date >= datetime.fromisoformat(request.args['date_from']))
    if request.args.get('date_to'):
        q = q.filter(Expense.date <= datetime.fromisoformat(request.args['date_to']))
    if request.args.get('category'):
        q = q.filter(Expense.category == request.args['category'])
    return q.order_by(Expense.date.desc()).all()


@export_bp.route('/csv', methods=['GET'])
@jwt_required()
def export_csv():
    uid = _get_user_id()
    expenses = _get_filtered_expenses(uid)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['Date', 'Merchant', 'Category', 'Amount', 'Currency', 'Tax', 'Notes'])
    for exp in expenses:
        writer.writerow([
            exp.date.strftime('%Y-%m-%d'),
            exp.merchant_name,
            exp.category,
            exp.amount,
            exp.currency,
            exp.tax_amount or 0,
            exp.notes or ''
        ])

    output.seek(0)
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'text/csv'
    response.headers['Content-Disposition'] = 'attachment; filename=expenses.csv'
    return response


@export_bp.route('/excel', methods=['GET'])
@jwt_required()
def export_excel():
    uid = _get_user_id()
    expenses = _get_filtered_expenses(uid)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Expenses'

        # Header row
        headers = ['Date', 'Merchant', 'Category', 'Amount', 'Currency', 'Tax', 'Notes']
        header_fill = PatternFill('solid', fgColor='1a1a2e')
        for col, hdr in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=hdr)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row, exp in enumerate(expenses, 2):
            ws.cell(row=row, column=1, value=exp.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=exp.merchant_name)
            ws.cell(row=row, column=3, value=exp.category)
            ws.cell(row=row, column=4, value=exp.amount)
            ws.cell(row=row, column=5, value=exp.currency)
            ws.cell(row=row, column=6, value=exp.tax_amount or 0)
            ws.cell(row=row, column=7, value=exp.notes or '')

        # Summary sheet
        ws2 = wb.create_sheet('Summary')
        from collections import defaultdict
        cat_totals = defaultdict(float)
        for exp in expenses:
            cat_totals[exp.category] += exp.amount
        ws2.cell(1, 1, 'Category').font = Font(bold=True)
        ws2.cell(1, 2, 'Total').font = Font(bold=True)
        for i, (cat, total) in enumerate(sorted(cat_totals.items()), 2):
            ws2.cell(i, 1, cat)
            ws2.cell(i, 2, round(total, 3))

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='expenses.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except ImportError:
        return jsonify({'error': 'openpyxl not installed'}), 500


@export_bp.route('/pdf', methods=['GET'])
@jwt_required()
def export_pdf():
    uid = _get_user_id()
    expenses = _get_filtered_expenses(uid)

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph('Expense Report', styles['Title']))
        elements.append(Paragraph(f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}', styles['Normal']))
        elements.append(Spacer(1, 20))

        data = [['Date', 'Merchant', 'Category', 'Amount', 'Currency']]
        for exp in expenses:
            data.append([
                exp.date.strftime('%Y-%m-%d'),
                exp.merchant_name[:30],
                exp.category,
                f'{exp.amount:.3f}',
                exp.currency
            ])

        table = Table(data, colWidths=[80, 160, 100, 70, 60])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1a2e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f4ff')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
        ]))
        elements.append(table)

        # Total
        total = sum(e.amount for e in expenses)
        elements.append(Spacer(1, 20))
        elements.append(Paragraph(f'<b>Total: {total:.3f} KWD</b>', styles['Normal']))

        doc.build(elements)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='expenses.pdf', mimetype='application/pdf')
    except ImportError:
        return jsonify({'error': 'reportlab not installed'}), 500
